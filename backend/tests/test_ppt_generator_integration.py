"""Integration tests for PPT Generator that create actual presentations."""
import pytest
import os
import sys
from datetime import datetime
from io import BytesIO
from pptx import Presentation

# Add parent directory to path for imports
sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), '..')))

from ppt_generator.generator import PPTGenerator
from orchestrator.models import (
    PresentationContent,
    TitleSlide,
    AgendaSlide,
    SlideContent,
    SummarySlide
)


class TestPPTGeneratorIntegration:
    """Integration tests for PPTGenerator class."""
    
    @pytest.fixture
    def sample_content(self):
        """Create sample presentation content."""
        return PresentationContent(
            title=TitleSlide(
                main_title="AI in Healthcare",
                subtitle="Transforming Patient Care"
            ),
            agenda=AgendaSlide(
                items=[
                    "Introduction to AI",
                    "Current Applications",
                    "Future Trends"
                ]
            ),
            slides=[
                SlideContent(
                    title="Introduction to AI",
                    content=[
                        "Artificial Intelligence is transforming healthcare",
                        "Machine learning enables better diagnostics",
                        "AI assists in treatment planning",
                        "Reduces healthcare costs",
                        "Improves patient outcomes"
                    ],
                    notes="This slide introduces the concept of AI in healthcare",
                    layout="bullet_list",
                    slide_type="content_slide"
                ),
                SlideContent(
                    title="Current Applications",
                    content=[
                        "Medical imaging analysis",
                        "Drug discovery",
                        "Patient monitoring",
                        "Administrative automation",
                        "Predictive analytics"
                    ],
                    notes="Examples of current AI applications",
                    layout="bullet_list",
                    slide_type="content_slide"
                )
            ],
            summary=SummarySlide(
                title="Key Takeaways",
                takeaways=[
                    "AI is revolutionizing healthcare",
                    "Multiple applications already in use",
                    "Future holds even more promise"
                ]
            )
        )
    
    def test_create_presentation_basic(self, sample_content):
        """Test basic presentation creation."""
        generator = PPTGenerator()
        
        result = generator.create_presentation(sample_content)
        
        # Should return bytes
        assert isinstance(result, bytes)
        assert len(result) > 0
        
        # Should be valid .pptx file
        prs = Presentation(BytesIO(result))
        
        # Should have correct number of slides:
        # 1 title + 1 agenda + 2 content + 1 summary = 5 slides
        assert len(prs.slides) == 5
    
    def test_create_presentation_slide_content(self, sample_content):
        """Test presentation has correct slide content."""
        generator = PPTGenerator()
        result = generator.create_presentation(sample_content)
        
        prs = Presentation(BytesIO(result))
        
        # Check title slide
        title_slide = prs.slides[0]
        assert sample_content.title.main_title in title_slide.shapes.title.text
        
        # Check agenda slide
        agenda_slide = prs.slides[1]
        assert "Agenda" in agenda_slide.shapes.title.text
        
        # Check content slides
        content_slide_1 = prs.slides[2]
        assert sample_content.slides[0].title in content_slide_1.shapes.title.text
        
        # Check summary slide
        summary_slide = prs.slides[4]
        assert sample_content.summary.title in summary_slide.shapes.title.text
    
    def test_create_presentation_with_many_slides(self):
        """Test presentation creation with many content slides."""
        content = PresentationContent(
            title=TitleSlide(main_title="Test Presentation"),
            agenda=AgendaSlide(items=["Topic 1", "Topic 2", "Topic 3"]),
            slides=[
                SlideContent(
                    title=f"Slide {i}",
                    content=[f"Point {j}" for j in range(1, 6)],
                    layout="bullet_list",
                    slide_type="content_slide"
                )
                for i in range(1, 11)  # 10 content slides
            ],
            summary=SummarySlide(
                title="Summary",
                takeaways=["Point 1", "Point 2", "Point 3"]
            )
        )
        
        generator = PPTGenerator()
        result = generator.create_presentation(content)
        
        prs = Presentation(BytesIO(result))
        
        # Should have 1 title + 1 agenda + 10 content + 1 summary = 13 slides
        assert len(prs.slides) == 13
    
    def test_create_presentation_without_subtitle(self):
        """Test presentation creation without subtitle."""
        content = PresentationContent(
            title=TitleSlide(main_title="Test Presentation", subtitle=None),
            agenda=AgendaSlide(items=["Topic 1"]),
            slides=[
                SlideContent(
                    title="Slide 1",
                    content=["Point 1", "Point 2", "Point 3", "Point 4", "Point 5"],
                    layout="bullet_list",
                    slide_type="content_slide"
                )
            ],
            summary=SummarySlide(title="Summary", takeaways=["Point 1"])
        )
        
        generator = PPTGenerator()
        result = generator.create_presentation(content)
        
        # Should succeed without subtitle
        assert isinstance(result, bytes)
        assert len(result) > 0
    
    def test_create_presentation_with_speaker_notes(self, sample_content):
        """Test presentation includes speaker notes."""
        generator = PPTGenerator()
        result = generator.create_presentation(sample_content)
        
        prs = Presentation(BytesIO(result))
        
        # Check first content slide has notes
        content_slide = prs.slides[2]
        notes_text = content_slide.notes_slide.notes_text_frame.text
        
        assert sample_content.slides[0].notes in notes_text
    
    def test_presentation_consistent_styling(self, sample_content):
        """Test that all slides use consistent styling."""
        generator = PPTGenerator()
        result = generator.create_presentation(sample_content)
        
        prs = Presentation(BytesIO(result))
        
        # Check that title slides have consistent formatting
        # This is a basic check - in a real scenario we'd check colors, fonts, etc.
        for i in range(1, len(prs.slides)):
            slide = prs.slides[i]
            # All slides should have a title
            assert slide.shapes.title is not None
