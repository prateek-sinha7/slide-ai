"""Unit tests for PPT Generator."""
import pytest
import os
import sys
from datetime import datetime
from io import BytesIO
from pptx import Presentation

# Add parent directory to path for imports
sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), '..')))

from ppt_generator.generator import (
    PPTGenerator,
    generate_filename,
    sanitize_topic,
    FileCreationError,
    TemplateNotFoundError
)
from orchestrator.models import (
    PresentationContent,
    TitleSlide,
    AgendaSlide,
    SlideContent,
    SummarySlide
)


class TestFilenameGeneration:
    """Tests for filename generation and sanitization."""
    
    def test_generate_filename_basic(self):
        """Test basic filename generation."""
        topic = "AI in Healthcare"
        timestamp = datetime(2024, 1, 15, 14, 30, 45)
        
        filename = generate_filename(topic, timestamp)
        
        assert filename == "AI_in_Healthcare_20240115_143045.pptx"
    
    def test_generate_filename_with_special_chars(self):
        """Test filename generation with special characters."""
        topic = "Test: Topic/With\\Special*Chars?"
        timestamp = datetime(2024, 1, 15, 14, 30, 45)
        
        filename = generate_filename(topic, timestamp)
        
        # Should remove all special characters
        assert "/" not in filename
        assert "\\" not in filename
        assert ":" not in filename
        assert "*" not in filename
        assert "?" not in filename
        assert filename.endswith(".pptx")
    
    def test_generate_filename_with_path_separators(self):
        """Test filename generation removes path separators."""
        topic = "../../../etc/passwd"
        timestamp = datetime(2024, 1, 15, 14, 30, 45)
        
        filename = generate_filename(topic, timestamp)
        
        # Should not contain path separators
        assert "/" not in filename
        assert ".." not in filename
        assert filename.endswith(".pptx")
    
    def test_generate_filename_empty_topic(self):
        """Test filename generation with empty topic."""
        topic = ""
        timestamp = datetime(2024, 1, 15, 14, 30, 45)
        
        filename = generate_filename(topic, timestamp)
        
        # Should use default name
        assert filename == "presentation_20240115_143045.pptx"
    
    def test_generate_filename_only_special_chars(self):
        """Test filename generation with only special characters."""
        topic = "***///:::"
        timestamp = datetime(2024, 1, 15, 14, 30, 45)
        
        filename = generate_filename(topic, timestamp)
        
        # Should use default name
        assert filename == "presentation_20240115_143045.pptx"
    
    def test_generate_filename_long_topic(self):
        """Test filename generation with very long topic."""
        topic = "A" * 200
        timestamp = datetime(2024, 1, 15, 14, 30, 45)
        
        filename = generate_filename(topic, timestamp)
        
        # Should be truncated to reasonable length
        assert len(filename) < 150
        assert filename.endswith(".pptx")
    
    def test_sanitize_topic_basic(self):
        """Test basic topic sanitization."""
        topic = "AI in Healthcare"
        result = sanitize_topic(topic)
        assert result == "AI in Healthcare"
    
    def test_sanitize_topic_special_chars(self):
        """Test topic sanitization with special characters."""
        topic = "Test: Topic/With\\Special*Chars?"
        result = sanitize_topic(topic)
        
        # Should remove special characters but keep spaces
        assert "/" not in result
        assert "\\" not in result
        assert ":" not in result
        assert "*" not in result
        assert "?" not in result
    
    def test_sanitize_topic_multiple_spaces(self):
        """Test topic sanitization collapses multiple spaces."""
        topic = "AI    in     Healthcare"
        result = sanitize_topic(topic)
        assert result == "AI in Healthcare"
    
    def test_sanitize_topic_leading_trailing_spaces(self):
        """Test topic sanitization removes leading/trailing spaces."""
        topic = "   AI in Healthcare   "
        result = sanitize_topic(topic)
        assert result == "AI in Healthcare"


class TestPPTGenerator:
    """Tests for PPTGenerator class."""
    
    @pytest.fixture
    def sample_content(self):
        """Create sample presentation content."""
        return PresentationContent(
            title=TitleSlide(
                main_title="AI in Healthcare",
                subtitle="Transforming Patient Care"
            ),
            agenda=AgendaSlide(
                items=[
                    "Introduction to AI",
                    "Current Applications",
                    "Future Trends"
                ]
            ),
            slides=[
                SlideContent(
                    title="Introduction to AI",
                    content=[
                        "Artificial Intelligence is transforming healthcare",
                        "Machine learning enables better diagnostics",
                        "AI assists in treatment planning",
                        "Reduces healthcare costs",
                        "Improves patient outcomes"
                    ],
                    notes="This slide introduces the concept of AI in healthcare",
                    layout="bullet_list",
                    slide_type="content_slide"
                ),
                SlideContent(
                    title="Current Applications",
                    content=[
                        "Medical imaging analysis",
                        "Drug discovery",
                        "Patient monitoring",
                        "Administrative automation",
                        "Predictive analytics"
                    ],
                    notes="Examples of current AI applications",
                    layout="bullet_list",
                    slide_type="content_slide"
                )
            ],
            summary=SummarySlide(
                title="Key Takeaways",
                takeaways=[
                    "AI is revolutionizing healthcare",
                    "Multiple applications already in use",
                    "Future holds even more promise"
                ]
            )
        )
    
    def test_create_presentation_basic(self, sample_content):
        """Test basic presentation creation."""
        generator = PPTGenerator()
        
        result = generator.create_presentation(sample_content)
        
        # Should return bytes
        assert isinstance(result, bytes)
        assert len(result) > 0
        
        # Should be valid .pptx file
        prs = Presentation(BytesIO(result))
        
        # Should have correct number of slides:
        # 1 title + 1 agenda + 2 content + 1 summary = 5 slides
        assert len(prs.slides) == 5
    
    def test_create_presentation_slide_content(self, sample_content):
        """Test presentation has correct slide content."""
        generator = PPTGenerator()
        result = generator.create_presentation(sample_content)
        
        prs = Presentation(BytesIO(result))
        
        # Check title slide
        title_slide = prs.slides[0]
        assert sample_content.title.main_title in title_slide.shapes.title.text
        
        # Check agenda slide
        agenda_slide = prs.slides[1]
        assert "Agenda" in agenda_slide.shapes.title.text
        
        # Check content slides
        content_slide_1 = prs.slides[2]
        assert sample_content.slides[0].title in content_slide_1.shapes.title.text
        
        # Check summary slide
        summary_slide = prs.slides[4]
        assert sample_content.summary.title in summary_slide.shapes.title.text
    
    def test_create_presentation_with_many_slides(self):
        """Test presentation creation with many content slides."""
        content = PresentationContent(
            title=TitleSlide(main_title="Test Presentation"),
            agenda=AgendaSlide(items=["Topic 1", "Topic 2", "Topic 3"]),
            slides=[
                SlideContent(
                    title=f"Slide {i}",
                    content=[f"Point {j}" for j in range(1, 6)],
                    layout="bullet_list",
                    slide_type="content_slide"
                )
                for i in range(1, 11)  # 10 content slides
            ],
            summary=SummarySlide(
                title="Summary",
                takeaways=["Point 1", "Point 2", "Point 3"]
            )
        )
        
        generator = PPTGenerator()
        result = generator.create_presentation(content)
        
        prs = Presentation(BytesIO(result))
        
        # Should have 1 title + 1 agenda + 10 content + 1 summary = 13 slides
        assert len(prs.slides) == 13
    
    def test_create_presentation_without_subtitle(self):
        """Test presentation creation without subtitle."""
        content = PresentationContent(
            title=TitleSlide(main_title="Test Presentation", subtitle=None),
            agenda=AgendaSlide(items=["Topic 1"]),
            slides=[
                SlideContent(
                    title="Slide 1",
                    content=["Point 1", "Point 2", "Point 3", "Point 4", "Point 5"],
                    layout="bullet_list",
                    slide_type="content_slide"
                )
            ],
            summary=SummarySlide(title="Summary", takeaways=["Point 1"])
        )
        
        generator = PPTGenerator()
        result = generator.create_presentation(content)
        
        # Should succeed without subtitle
        assert isinstance(result, bytes)
        assert len(result) > 0
    
    def test_create_presentation_with_speaker_notes(self, sample_content):
        """Test presentation includes speaker notes."""
        generator = PPTGenerator()
        result = generator.create_presentation(sample_content)
        
        prs = Presentation(BytesIO(result))
        
        # Check first content slide has notes
        content_slide = prs.slides[2]
        notes_text = content_slide.notes_slide.notes_text_frame.text
        
        assert sample_content.slides[0].notes in notes_text
    
    def test_create_presentation_template_not_found(self, sample_content):
        """Test error handling when template not found."""
        generator = PPTGenerator(template_dir="/nonexistent/path")
        
        # Should still work with blank presentation
        result = generator.create_presentation(sample_content, template="nonexistent")
        
        # Should return valid presentation
        assert isinstance(result, bytes)
        assert len(result) > 0
    
    def test_create_presentation_invalid_content(self):
        """Test error handling with invalid content."""
        generator = PPTGenerator()
        
        # This should raise an error due to invalid content structure
        with pytest.raises(Exception):
            generator.create_presentation(None)
    
    def test_presentation_consistent_styling(self, sample_content):
        """Test that all slides use consistent styling."""
        generator = PPTGenerator()
        result = generator.create_presentation(sample_content)
        
        prs = Presentation(BytesIO(result))
        
        # Check that title slides have consistent formatting
        # This is a basic check - in a real scenario we'd check colors, fonts, etc.
        for i in range(1, len(prs.slides)):
            slide = prs.slides[i]
            # All slides should have a title
            assert slide.shapes.title is not None


class TestPPTGeneratorErrors:
    """Tests for PPT Generator error handling."""
    
    def test_file_creation_error_handling(self):
        """Test that file creation errors are properly wrapped."""
        generator = PPTGenerator()
        
        # Create invalid content that will cause an error
        invalid_content = "not a PresentationContent object"
        
        with pytest.raises((FileCreationError, AttributeError, TypeError)):
            generator.create_presentation(invalid_content)
